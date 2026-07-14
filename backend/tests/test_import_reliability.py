"""Regression coverage for resilient long-document AI imports."""

from backend.imports.image_extractor import ImagePayload


def _question(text: str) -> dict:
    return {
        "type": "fill_blank",
        "question": text,
        "answer": "答案",
        "analysis": "解析",
    }


def test_docx_extraction_preserves_paragraph_table_order(tmp_path):
    """Table questions must remain between the surrounding Word paragraphs."""
    from docx import Document
    from backend.imports import import_orchestrator

    document = Document()
    document.add_paragraph("1. First question")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A. First option"
    table.cell(0, 1).text = "B. Second option"
    document.add_paragraph("Answer: A")
    source = tmp_path / "ordered.docx"
    document.save(source)

    text, warnings = import_orchestrator.extract_text_and_warnings(str(source))

    assert warnings == []
    assert text.index("1. First question") < text.index("A. First option") < text.index("Answer: A")


def test_rule_parser_keeps_subquestions_and_steps_inside_their_top_level_question(monkeypatch):
    """Only continuous ``number + punctuation`` lines start a new question."""
    from backend.imports import import_orchestrator

    text = """一、单项选择题
1. First question
A. Alpha
B. Beta
答案：A
2. Second question
A. Alpha
B. Beta
答案：B
二、多项选择题
3. Third question
A. Alpha
B. Beta
C. Gamma
答案：A、C
三、判断题
4. A true false question
答案：正确
四、名词解释
5. Define a term
答案：A definition
五、简答与计算题
6. Solve the following
（1）show the calculation
① first answer step
答案：A complete answer
六、讨论题
7. Discuss the result
答案：A discussion answer"""

    monkeypatch.setattr(
        import_orchestrator,
        "call_ai_parse",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("regular document must not call AI")),
    )

    questions, warnings, timing = import_orchestrator.preview_import_from_text(text)

    assert warnings == []
    assert timing["rule_based"] is True
    assert timing["is_complete"] is True
    assert [question["line_number"] for question in questions] == list(range(1, 8))
    assert [question["type"] for question in questions] == [
        "single_choice",
        "single_choice",
        "multiple_choice",
        "true_false",
        "short_answer",
        "short_answer",
        "short_answer",
    ]
    assert "（1）show the calculation" in questions[5]["question"]
    assert "① first answer step" in questions[5]["question"]


def test_rule_parser_handles_a_94_question_six_section_bank_without_ai(monkeypatch):
    """The known 38/8/30/9/6/3 structure must remain 94 main questions."""
    from collections import Counter
    from backend.imports import import_orchestrator

    sections = [
        ("一、单项选择题", 38, "choice"),
        ("二、多项选择题", 8, "choice"),
        ("三、判断题", 30, "boolean"),
        ("四、名词解释", 9, "short"),
        ("五、简答与计算题", 6, "short"),
        ("六、讨论题", 3, "short"),
    ]
    lines: list[str] = []
    number = 1
    for heading, count, kind in sections:
        lines.append(heading)
        for _ in range(count):
            lines.append(f"{number}. Simulated question {number}")
            if kind == "choice":
                lines.extend(["A. First option", "B. Second option", "答案：A"])
            elif kind == "boolean":
                lines.append("答案：正确")
            else:
                lines.extend(["（1）subquestion stays in this item", "① answer step stays in this item", "答案：完整答案"])
            number += 1

    monkeypatch.setattr(
        import_orchestrator,
        "call_ai_parse",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("regular 94-question bank must not call AI")),
    )
    questions, warnings, timing = import_orchestrator.preview_import_from_text("\n".join(lines))

    assert warnings == []
    assert timing["is_complete"] is True
    assert [question["line_number"] for question in questions] == list(range(1, 95))
    assert Counter(question["type"] for question in questions) == {
        "single_choice": 38,
        "multiple_choice": 8,
        "true_false": 30,
        "short_answer": 18,
    }


def test_rule_parser_marks_non_continuous_numbering_partial_without_ai_fallback():
    from backend.imports import import_orchestrator

    text = """一、判断题
1. First statement
答案：正确
3. Missing-number statement
答案：错误"""

    questions, warnings, timing = import_orchestrator.parse_rule_based_question_document(text)

    assert [question["line_number"] for question in questions] == [1]
    assert timing["is_complete"] is False
    assert any("题号 3" in warning for warning in warnings)


def test_nested_ai_response_questions_are_accepted():
    from backend.imports import import_orchestrator

    assert import_orchestrator.question_items_from_parsed_json({"data": {"questions": [_question("nested")]}}) == [
        _question("nested")
    ]


def test_rule_parser_does_not_duplicate_embedded_images(monkeypatch):
    from backend.imports import import_orchestrator

    text = """一、判断题
1. 如下图所示，判断结论是否正确。
答案：正确"""
    monkeypatch.setattr(
        import_orchestrator,
        "call_ai_parse_multimodal",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("must not create detached image questions")),
    )

    questions, warnings, timing = import_orchestrator.preview_import_from_file_content(
        text,
        [ImagePayload(data=b"image", mime_type="image/png", source="diagram")],
    )

    assert len(questions) == 1
    assert timing["is_complete"] is False
    assert timing["failed_chunks"] == 1
    assert any("题图" in warning for warning in warnings)


def test_chunk_document_text_splits_one_large_numbered_paragraph(monkeypatch):
    """A Word document can store many numbered questions in one paragraph."""
    from backend.imports import import_orchestrator

    monkeypatch.setattr(import_orchestrator, "AI_CHUNK_SIZE", 80)
    text = " ".join(
        f"{index}. 这是第 {index} 道很长的模拟题，包含足够多的文字以触发分块。答案：答案。"
        for index in range(1, 7)
    )

    chunks, _ = import_orchestrator.chunk_document_text(text)

    assert len(chunks) >= 3
    assert all(len(chunk) <= 120 for chunk in chunks)
    assert "1." in chunks[0]
    assert "6." in chunks[-1]


def test_chunk_document_text_caps_short_numbered_questions_per_ai_request(monkeypatch):
    """Short Word paragraphs must not let one AI request contain dozens of questions."""
    from backend.imports import import_orchestrator

    monkeypatch.setattr(import_orchestrator, "AI_CHUNK_SIZE", 10_000)
    text = "\n".join(f"{index}. Simulated question {index}" for index in range(1, 15))

    chunks, _ = import_orchestrator.chunk_document_text(text)

    assert len(chunks) == 3
    assert all(import_orchestrator.count_numbered_question_blocks(chunk) <= 6 for chunk in chunks)
    assert "1." in chunks[0]
    assert "14." in chunks[-1]


def test_incomplete_numbered_chunk_is_retried_in_smaller_batches(monkeypatch):
    """A partial JSON response must be retried instead of silently importing a subset."""
    from backend.imports import import_orchestrator

    text = "\n".join(f"{index}. Simulated question {index}" for index in range(1, 7))
    calls: list[int] = []

    def fake_parse(chunk: str, _index: int, expected_question_count: int = 0):
        calls.append(expected_question_count)
        numbers = [int(match.group(1)) for match in __import__("re").finditer(r"(?m)^(\d+)\.", chunk)]
        if expected_question_count == 6:
            return [_question("question 1")], []
        return [_question(f"question {number}") for number in numbers], []

    monkeypatch.setattr(import_orchestrator, "call_ai_parse_chunk", fake_parse)

    items, warnings = import_orchestrator.parse_chunk_with_coverage_retry(text, 0)

    assert [item["question"] for item in items] == [f"question {index}" for index in range(1, 7)]
    assert calls == [6, 2, 2, 2]
    assert any("自动拆分重试" in warning for warning in warnings)


def test_numbered_document_over_batch_size_processes_every_chunk(monkeypatch):
    """A per-batch limit must not reject or drop later document chunks."""
    from backend.imports import import_orchestrator

    monkeypatch.setattr(import_orchestrator, "OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(import_orchestrator, "AI_BATCH_SIZE", 1)
    text = "\n".join(f"{index}. Simulated question {index}" for index in range(1, 13))

    def fake_parse(chunk: str, _index: int, expected_question_count: int = 0):
        numbers = [int(match.group(1)) for match in __import__("re").finditer(r"(?m)^(\d+)\.", chunk)]
        assert len(numbers) == expected_question_count
        return [_question(f"question {number}") for number in numbers], []

    monkeypatch.setattr(import_orchestrator, "call_ai_parse_chunk", fake_parse)

    questions, warnings, timing = import_orchestrator.call_ai_parse(text)

    assert len(questions) == 12
    assert warnings == []
    assert timing["chunks"] == 2
    assert timing["batches"] == 2
    assert timing["completed_chunks"] == 2
    assert timing["is_complete"] is True


def test_long_document_reports_persistable_progress_after_each_chunk(monkeypatch):
    """Progress checkpoints must cover every chunk, not only the final result."""
    from backend.imports import import_orchestrator

    monkeypatch.setattr(import_orchestrator, "OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(import_orchestrator, "AI_BATCH_SIZE", 1)
    text = "\n".join(f"{index}. Simulated question {index}" for index in range(1, 13))
    checkpoints: list[tuple[int, int, int]] = []

    def fake_parse(chunk: str, _index: int, _expected_question_count: int = 0):
        numbers = [int(match.group(1)) for match in __import__("re").finditer(r"(?m)^(\d+)\.", chunk)]
        return [_question(f"question {number}") for number in numbers], []

    monkeypatch.setattr(import_orchestrator, "call_ai_parse_chunk", fake_parse)

    import_orchestrator.call_ai_parse(
        text,
        on_progress=lambda questions, _warnings, timing: checkpoints.append(
            (len(questions), timing["completed_chunks"], timing["chunks"])
        ),
    )

    assert checkpoints[0] == (0, 0, 2)
    assert checkpoints[-1] == (12, 2, 2)
    assert len(checkpoints) == 3


def test_transient_chunk_failure_retries_without_restarting_document(monkeypatch):
    """A temporary upstream timeout should retry only the affected chunk."""
    from fastapi import HTTPException
    from backend.imports import import_orchestrator

    attempts = 0

    def fake_parse(_chunk: str, _index: int):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise HTTPException(status_code=504, detail="timeout")
        return [_question("recovered question")], []

    monkeypatch.setattr(import_orchestrator, "parse_chunk_with_coverage_retry", fake_parse)

    questions, warnings = import_orchestrator.parse_chunk_with_transient_retry("1. Question", 0)

    assert attempts == 2
    assert [item["question"] for item in questions] == ["recovered question"]
    assert warnings == []


def test_non_object_ai_items_are_skipped_without_crashing(monkeypatch):
    """A scalar inside an AI JSON array must not crash deduplication."""
    from backend.imports import import_orchestrator

    monkeypatch.setattr(import_orchestrator, "OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(import_orchestrator, "AI_BATCH_SIZE", 1)
    monkeypatch.setattr(
        import_orchestrator,
        "call_ai_parse_chunk",
        lambda _chunk, _index, _expected_question_count=0: (["not-a-question", _question("valid question")], []),
    )

    questions, warnings, timing = import_orchestrator.call_ai_parse("1. Simulated question")

    assert [item["question"] for item in questions] == ["valid question"]
    assert timing["is_complete"] is False
    assert any("无效题目记录" in warning for warning in warnings)


def test_invalid_ai_field_types_become_partial_preview(monkeypatch):
    """Nested AI fields must be rejected without crashing task aggregation."""
    from backend.imports import import_orchestrator

    monkeypatch.setattr(import_orchestrator, "OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(import_orchestrator, "AI_BATCH_SIZE", 1)
    malformed = {
        "type": "fill_blank",
        "question": ["not text"],
        "answer": "answer",
        "analysis": {"nested": "value"},
        "subject": ["nested"],
    }
    monkeypatch.setattr(
        import_orchestrator,
        "call_ai_parse_chunk",
        lambda _chunk, _index, _expected_question_count=0: ([malformed, _question("valid question")], []),
    )

    questions, warnings, timing = import_orchestrator.call_ai_parse("1. Simulated question")

    assert [item["question"] for item in questions] == ["valid question"]
    assert timing["is_complete"] is False
    assert any("字段格式异常" in warning for warning in warnings)


def test_file_content_combines_chunked_text_and_image_questions(monkeypatch):
    """Embedded images must not make a long text document skip text parsing."""
    from backend.imports import import_orchestrator

    text_questions = [_question("文本题一"), _question("文本题二")]
    image_questions = [_question("图片题一")]
    monkeypatch.setattr(
        import_orchestrator,
        "call_ai_parse",
        lambda _text: (text_questions, ["文本解析提示"], {"chunks": 2, "ai_ms": 20, "ai_chunks": [10, 10]}),
    )
    monkeypatch.setattr(
        import_orchestrator,
        "call_ai_parse_multimodal",
        lambda _text, _images: (image_questions, ["图片解析提示"], {"chunks": 1, "ai_ms": 10, "ai_chunks": [10]}),
    )

    questions, warnings, timing = import_orchestrator.preview_import_from_file_content(
        "1. 文本题一\n2. 文本题二",
        [ImagePayload(data=b"image", mime_type="image/png", source="题图")],
    )

    assert [item["question"] for item in questions] == ["文本题一", "文本题二", "图片题一"]
    assert warnings == ["文本解析提示", "图片解析提示"]
    assert timing["chunks"] == 3
